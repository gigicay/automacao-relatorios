from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter


def gerar_excel(df):
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo executivo"
    analise = wb.create_sheet("Análises")
    dados = wb.create_sheet("Dados tratados")

    navy = "17243A"
    blue = "2F6FEB"
    light_blue = "EAF1FF"
    dark_text = "243044"
    muted = "6B7890"
    white = "FFFFFF"
    light = "F4F7FB"
    line = "D9E0EA"
    green = "16A085"

    thin = Side(style="thin", color=line)

    faturamento = float(df["valor_total"].sum())
    vendas = int(len(df))
    clientes = int(df["cliente"].nunique())
    ticket = faturamento / vendas if vendas else 0

    # ===== RESUMO EXECUTIVO =====
    resumo.sheet_view.showGridLines = False
    resumo.merge_cells("A1:H1")
    resumo["A1"] = "RELATÓRIO DE VENDAS"
    resumo["A1"].font = Font(size=20, bold=True, color=white)
    resumo["A1"].fill = PatternFill("solid", fgColor=navy)
    resumo["A1"].alignment = Alignment(vertical="center")
    resumo.row_dimensions[1].height = 34

    resumo.merge_cells("A2:H2")
    resumo["A2"] = "Resumo executivo dos dados tratados"
    resumo["A2"].font = Font(size=10, color=muted, italic=True)
    resumo["A2"].alignment = Alignment(vertical="center")
    resumo.row_dimensions[2].height = 23

    cards = [
        ("A4:B4", "A5:B6", "Faturamento total", faturamento, 'R$ #,##0.00'),
        ("C4:D4", "C5:D6", "Total de vendas", vendas, '#,##0'),
        ("E4:F4", "E5:F6", "Clientes", clientes, '#,##0'),
        ("G4:H4", "G5:H6", "Ticket médio", ticket, 'R$ #,##0.00'),
    ]
    for label_range, value_range, label, value, fmt in cards:
        resumo.merge_cells(label_range)
        resumo.merge_cells(value_range)
        lc = resumo[label_range.split(":")[0]]
        vc = resumo[value_range.split(":")[0]]
        lc.value = label
        lc.font = Font(size=10, bold=True, color=muted)
        lc.fill = PatternFill("solid", fgColor=light)
        vc.value = value
        vc.number_format = fmt
        vc.font = Font(size=17, bold=True, color=navy)
        vc.fill = PatternFill("solid", fgColor=light)
        lc.alignment = vc.alignment = Alignment(vertical="center")
        for row in resumo[label_range]:
            for c in row:
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in resumo[value_range]:
            for c in row:
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Product summary table
    resumo["A9"] = "TOP PRODUTOS POR FATURAMENTO"
    resumo["A9"].font = Font(size=12, bold=True, color=navy)
    prod = df.groupby("produto", as_index=False)["valor_total"].sum().sort_values("valor_total", ascending=False)
    resumo["A10"] = "Produto"
    resumo["B10"] = "Faturamento"
    for c in resumo[10][:2]:
        c.font = Font(bold=True, color=white)
        c.fill = PatternFill("solid", fgColor=blue)

    for i, row in enumerate(prod.head(10).itertuples(index=False), 11):
        resumo.cell(i, 1).value = row.produto
        resumo.cell(i, 2).value = float(row.valor_total)
        resumo.cell(i, 2).number_format = 'R$ #,##0.00'

    chart = BarChart()
    chart.title = "Top produtos"
    chart.style = 10
    chart.height = 7
    chart.width = 12
    data = Reference(resumo, min_col=2, min_row=10, max_row=10 + min(10, len(prod)))
    cats = Reference(resumo, min_col=1, min_row=11, max_row=10 + min(10, len(prod)))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    resumo.add_chart(chart, "D9")

    for col, width in {"A": 28, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18}.items():
        resumo.column_dimensions[col].width = width

    # ===== ANÁLISES =====
    analise.sheet_view.showGridLines = False
    analise.merge_cells("A1:H1")
    analise["A1"] = "ANÁLISES DE VENDAS"
    analise["A1"].font = Font(size=18, bold=True, color=white)
    analise["A1"].fill = PatternFill("solid", fgColor=navy)
    analise.row_dimensions[1].height = 32

    # Category
    cat = df.groupby("categoria", as_index=False)["valor_total"].sum().sort_values("valor_total", ascending=False)
    analise["A4"] = "FATURAMENTO POR CATEGORIA"
    analise["A4"].font = Font(size=12, bold=True, color=navy)
    analise["A5"] = "Categoria"
    analise["B5"] = "Faturamento"
    for c in analise[5][:2]:
        c.font = Font(bold=True, color=white)
        c.fill = PatternFill("solid", fgColor=blue)
    for i, row in enumerate(cat.itertuples(index=False), 6):
        analise.cell(i, 1).value = row.categoria
        analise.cell(i, 2).value = float(row.valor_total)
        analise.cell(i, 2).number_format = 'R$ #,##0.00'

    # Monthly
    analise["D4"] = "EVOLUÇÃO MENSAL"
    analise["D4"].font = Font(size=12, bold=True, color=navy)
    mensal = df.groupby("mes", as_index=False)["valor_total"].sum().sort_values("mes")
    analise["D5"] = "Mês"
    analise["E5"] = "Faturamento"
    for c in analise[5][3:5]:
        c.font = Font(bold=True, color=white)
        c.fill = PatternFill("solid", fgColor=blue)
    for i, row in enumerate(mensal.itertuples(index=False), 6):
        analise.cell(i, 4).value = row.mes
        analise.cell(i, 5).value = float(row.valor_total)
        analise.cell(i, 5).number_format = 'R$ #,##0.00'

    line_chart = LineChart()
    line_chart.title = "Evolução do faturamento"
    line_chart.style = 13
    line_chart.height = 7
    line_chart.width = 12
    data_ref = Reference(analise, min_col=5, min_row=5, max_row=5 + len(mensal))
    cats_ref = Reference(analise, min_col=4, min_row=6, max_row=5 + len(mensal))
    line_chart.add_data(data_ref, titles_from_data=True)
    line_chart.set_categories(cats_ref)
    analise.add_chart(line_chart, "G4")

    for col, width in {"A": 25, "B": 18, "C": 4, "D": 18, "E": 18, "F": 4, "G": 18, "H": 18}.items():
        analise.column_dimensions[col].width = width

    # ===== DADOS TRATADOS =====
    dados.sheet_view.showGridLines = False
    columns = list(df.columns)
    for idx, name in enumerate(columns, 1):
        cell = dados.cell(1, idx)
        cell.value = name.replace("_", " ").title()
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)

    for r, values in enumerate(df.itertuples(index=False, name=None), 2):
        for c, value in enumerate(values, 1):
            cell = dados.cell(r, c)
            cell.value = value
            cell.alignment = Alignment(vertical="center")
            name = columns[c - 1]
            if name == "data":
                cell.number_format = "dd/mm/yyyy"
            elif name in ("preco", "valor_total"):
                cell.number_format = 'R$ #,##0.00'

    dados.freeze_panes = "A2"
    dados.auto_filter.ref = dados.dimensions
    dados.row_dimensions[1].height = 25

    for c, name in enumerate(columns, 1):
        max_len = len(str(name)) + 3
        for r in range(2, min(dados.max_row, 101) + 1):
            max_len = max(max_len, len(str(dados.cell(r, c).value or "")))
        dados.column_dimensions[get_column_letter(c)].width = min(max(max_len, 12), 32)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
